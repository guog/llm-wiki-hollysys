#!/usr/bin/env python3
"""Generate tiny parser fixtures in a caller-provided raw directory."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def main() -> None:
    target = Path(sys.argv[1])
    target.mkdir(parents=True, exist_ok=True)

    (target / "sample.md").write_text(
        "# Alpha Handbook\n\nAlpha control limit is 42 units.\n\nSee the alpha concept for operating guidance.\n",
        encoding="utf-8",
    )
    (target / "sample.csv").write_text("name,value\nCSV_MARKER,42\n", encoding="utf-8")
    (target / "sample.html").write_text(
        "<html><body><h1>HTML_MARKER</h1><p>HTML knowledge.</p></body></html>", encoding="utf-8"
    )
    (target / "sample.json").write_text(json.dumps({"marker": "JSON_MARKER", "value": 42}), encoding="utf-8")
    (target / "sample.txt").write_text("TXT_MARKER plain knowledge.\n", encoding="utf-8")

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF_MARKER alpha PDF knowledge")
    pdf.save(target / "sample.pdf")
    pdf.close()

    document = Document()
    document.add_heading("DOCX_MARKER", level=1)
    document.add_paragraph("Alpha DOCX knowledge")
    document.save(target / "sample.docx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    text_box.text_frame.text = "PPTX_MARKER alpha presentation knowledge"
    presentation.save(target / "sample.pptx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "XLSX_MARKER"
    sheet["A2"] = "Limit"
    sheet["B2"] = 42
    sheet["C2"] = "=B2*2"
    hidden = workbook.create_sheet("Hidden Sheet")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "HIDDEN_MARKER"
    workbook.save(target / "sample.xlsx")


if __name__ == "__main__":
    main()
