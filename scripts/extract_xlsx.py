#!/usr/bin/env python3
"""Convert XLSX workbooks to source-grounded canonical Markdown."""

from __future__ import annotations

import argparse
import hashlib
from datetime import datetime, timezone
from pathlib import Path

import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def escape_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\\", "\\\\").replace("|", "\\|")
    return text.replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")


def workbook_to_markdown(source: Path, source_path: str) -> str:
    checksum = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, data_only=False, read_only=False)
    metadata = {
        "canonical_schema": 1,
        "source_path": source_path,
        "source_checksum": f"sha256:{checksum}",
        "source_type": "xlsx",
        "canonicalizer": "openpyxl",
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }
    lines = ["---", yaml.safe_dump(metadata, allow_unicode=True, sort_keys=False).strip(), "---", ""]
    lines.append(f"# Workbook: {source.name}")
    lines.append("")

    for sheet in workbook.worksheets:
        lines.extend(
            [
                f"## Sheet: {sheet.title}",
                "",
                f"- Dimensions: `{sheet.calculate_dimension()}`",
                f"- Visibility: `{sheet.sheet_state}`",
            ]
        )
        merged = [str(item) for item in sheet.merged_cells.ranges]
        lines.append(f"- Merged ranges: `{', '.join(merged) if merged else 'none'}`")
        lines.append("")

        non_empty_columns = [
            column
            for column in range(1, sheet.max_column + 1)
            if any(sheet.cell(row=row, column=column).value is not None for row in range(1, sheet.max_row + 1))
        ]
        if not non_empty_columns:
            lines.extend(["_Empty sheet._", ""])
            continue

        headers = ["Row"] + [get_column_letter(column) for column in non_empty_columns]
        lines.append("| " + " | ".join(headers) + " |")
        lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in range(1, sheet.max_row + 1):
            values = [sheet.cell(row=row, column=column).value for column in non_empty_columns]
            if all(value is None for value in values):
                continue
            rendered = [str(row)]
            for column, value in zip(non_empty_columns, values, strict=True):
                coordinate = f"{get_column_letter(column)}{row}"
                rendered.append(f"`{coordinate}` {escape_cell(value)}")
            lines.append("| " + " | ".join(rendered) + " |")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--source-path", required=True)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(workbook_to_markdown(args.source, args.source_path), encoding="utf-8")


if __name__ == "__main__":
    main()
